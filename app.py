import streamlit as st
from pypdf import PdfReader
import docx2txt
import re
import io
import pandas as pd
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# 1. Page Configuration & Professional Styling
st.set_page_config(page_title="Limitless Blue AI", page_icon="🔷", layout="wide")

st.markdown("""
    <style>
    .main { background-color: #0b111e; color: #ffffff; }
    h1 { color: #4e8cff; font-family: 'Helvetica Neue', sans-serif; font-weight: 700; }
    h3 { color: #94a3b8; font-weight: 400; }
    h4 { color: #f8fafc; border-bottom: 2px solid #1d4ed8; padding-bottom: 6px; margin-top: 15px; }
    .stButton>button { background-color: #1d4ed8; color: white; border-radius: 6px; width: 100%; height: 45px; font-weight: bold; border: none; }
    .metric-box { padding: 14px; background-color: #162238; border-radius: 8px; margin: 10px 0; border-left: 4px solid #4e8cff; font-size: 14px; line-height: 1.6; }
    .risk-box { padding: 14px; background-color: #2b161d; border-radius: 8px; margin: 10px 0; border-left: 4px solid #ef4444; font-size: 14px; line-height: 1.6; }
    </style>
    """, unsafe_allow_html=True)

# 2. Executive Header Layout
st.markdown("# 🔷 Limitless Blue | Enterprise Document Intelligence Engine")
st.markdown("### Secure, localized data extraction for corporate and legal compliance evaluation.")
st.markdown("---")

# 3. File Processing Module Supporting PDF, DOCX, and TXT
uploaded_file = st.file_uploader("Upload Corporate File (PDF, DOCX, or TXT)", type=["pdf", "docx", "txt"])

if uploaded_file is not None:
    text_content = ""
    file_type = uploaded_file.name.split(".")[-1].lower()
    
    try:
        if file_type == "pdf":
            reader = PdfReader(uploaded_file)
            for page in reader.pages:
                text_content += page.extract_text() or ""
        elif file_type == "docx":
            text_content = docx2txt.process(uploaded_file)
        elif file_type == "txt":
            text_content = uploaded_file.read().decode("utf-8")
            
        st.success(f"Successfully analyzed: {uploaded_file.name}")
        
        # --- LOCAL NLP EXTRACTION MATRICES ---
        sentences = re.split(r'(?<=[.!?])\s+', text_content)
        
        financial_insights = []
        risk_insights = []
        operational_insights = []
        
        for sentence in sentences:
            s_lower = sentence.lower()
            
            # Financial Data Mapping
            if any(k in s_lower for k in ["$", "revenue", "profit", "expenditure", "cost", "margin", "penalty"]):
                if len(sentence.strip()) > 15:
                    financial_insights.append(sentence.strip())
                    
            # Legal Liabilities Mapping
            if any(k in s_lower for k in ["liabilit", "breach", "damages", "nondisclosure", "nda", "comply", "compliance", "restrict", "penalty"]):
                if len(sentence.strip()) > 15:
                    risk_insights.append(sentence.strip())
                    
            # Operations Mapping
            if any(k in s_lower for k in ["protocol", "tenet", "must", "extract", "auth", "step", "schedul", "scrub"]):
                if len(sentence.strip()) > 15:
                    operational_insights.append(sentence.strip())

        # Eliminate duplication matrix
        financial_insights = list(set(financial_insights))[:5]
        risk_insights = list(set(risk_insights))[:5]
        operational_insights = list(set(operational_insights))[:5]

        # --- DYNAMIC DASHBOARD OUTPUT ---
        st.subheader("📋 Executive Intelligence Summary")
        
        col1, col2 = st.columns(2)
        
        with col1:
            st.markdown("#### Financial Metrics & Performance")
            if financial_insights:
                for item in financial_insights:
                    st.markdown(f"<div class='metric-box'><b>[FINANCIAL]</b> {item}</div>", unsafe_allow_html=True)
            else:
                st.info("No explicitly isolated financial metrics found.")
                
            st.markdown("#### Strategic & Operational Directives")
            if operational_insights:
                for item in operational_insights:
                    st.markdown(f"<div class='metric-box'><b>[OPERATIONS]</b> {item}</div>", unsafe_allow_html=True)
            else:
                st.info("No distinct operational steps detected.")

        with col2:
            st.markdown("#### Legal Liabilities & Risk Assessment")
            if risk_insights:
                for item in risk_insights:
                    st.markdown(f"<div class='risk-box'><b>[CRITICAL RISK]</b> {item}</div>", unsafe_allow_html=True)
            else:
                st.info("No critical legal compliance flags detected.")
                
            # --- FIXED DYNAMIC CHART ---
            st.markdown("#### Data Extraction Insights Density")
            metric_counts = {
                "Metrics": ["Financial Elements", "Operational Steps", "Risk Factors"],
                "Count": [len(financial_insights), len(operational_insights), len(risk_insights)]
            }
            chart_df = pd.DataFrame(metric_counts)
            st.bar_chart(data=chart_df, x="Metrics", y="Count", use_container_width=True)
        
        st.markdown("---")
        
        # --- SECURE MULTI-FORMAT EXPORT GENERATOR HUB ---
        st.subheader("📥 Export Intelligence Report Hub")
        format_choice = st.selectbox("Select Export Format Preference:", ["Portable Document Format (.pdf)", "Microsoft Word (.docx)", "Plain Text Document (.txt)"])
        
        # Text Document Builder (.txt)
        def build_txt_report():
            report = f"LIMITLESS BLUE EXECUTIVE INTEL RECAP: {uploaded_file.name.upper()}\n" + "="*50 + "\n\n"
            report += "[FINANCIAL MILESTONES]\n" + "\n".join([f"- {f}" for f in financial_insights]) + "\n\n"
            report += "[OPERATIONAL DIRECTIVES]\n" + "\n".join([f"- {o}" for o in operational_insights]) + "\n\n"
            report += "[CRITICAL COMPLIANCE RISKS]\n" + "\n".join([f"- {r}" for r in risk_insights]) + "\n"
            return report.encode('utf-8')

        # Microsoft Word Document Builder (.docx)
        def build_docx_report():
            doc = Document()
            doc.add_heading("Limitless Blue AI | Summary Intelligence Report", 0)
            doc.add_paragraph(f"Analyzed File Reference: {uploaded_file.name}")
            
            doc.add_heading("Financial Metrics & Performance Index", level=1)
            for item in financial_insights: doc.add_paragraph(item, style='List Bullet')
            
            doc.add_heading("Strategic & Operational Directives", level=1)
            for item in operational_insights: doc.add_paragraph(item, style='List Bullet')
            
            doc.add_heading("Legal Liabilities & Risk Assessment", level=1)
            for item in risk_insights: doc.add_paragraph(item, style='List Bullet')
            
            output_stream = io.BytesIO()
            doc.save(output_stream)
            return output_stream.getvalue()

        # Streamlined PDF Document Builder (.pdf)
        def build_pdf_report():
            output_stream = io.BytesIO()
            pdf_doc = SimpleDocTemplate(output_stream, pagesize=letter, title="Limitless Blue Intel Review")
            styles = getSampleStyleSheet()
            
            # Simple, standard spacing structure to avoid layout errors
            document_content = [
                Paragraph("<font size=22 color='#1d4ed8'><b>🔷 Limitless Blue Executive Report</b></font>", styles['Title']),
                Paragraph(f"<b>Extracted Audit Summary: {uploaded_file.name}</b>", styles['Normal']),
                Spacer(1, 15),
                Paragraph("<font size=14 color='#1e3a8a'><b>Financial Metrics & Performance</b></font>", styles['Heading2']),
                Spacer(1, 5)
            ]
            for item in financial_insights:
                document_content.append(Paragraph(f"• {item}", styles['Normal']))
                document_content.append(Spacer(1, 6))
                
            document_content.append(Spacer(1, 10))
            document_content.append(Paragraph("<font size=14 color='#1e3a8a'><b>Strategic & Operational Directives</b></font>", styles['Heading2']))
            document_content.append(Spacer(1, 5))
            for item in operational_insights:
                document_content.append(Paragraph(f"• {item}", styles['Normal']))
                document_content.append(Spacer(1, 6))
                
            document_content.append(Spacer(1, 10))
            document_content.append(Paragraph("<font size=14 color='#7f1d1d'><b>Legal Liabilities & Critical Compliance Risks</b></font>", styles['Heading2']))
            document_content.append(Spacer(1, 5))
            for item in risk_insights:
                document_content.append(Paragraph(f"• {item}", styles['Normal']))
                document_content.append(Spacer(1, 6))
                
            pdf_doc.build(document_content)
            return output_stream.getvalue()

        # Render active download streams cleanly matching your select boxes
        if "(.pdf)" in format_choice:
            st.download_button(label="📥 Download Summary Report (.pdf)", data=build_pdf_report(), file_name=f"Intel_Report_{uploaded_file.name}.pdf", mime="application/pdf")
        elif "(.docx)" in format_choice:
            st.download_button(label="📥 Download Summary Report (.docx)", data=build_docx_report(), file_name=f"Intel_Report_{uploaded_file.name}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        else:
